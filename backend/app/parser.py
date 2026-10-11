from __future__ import annotations
from io import BytesIO
import re

def clean(text: str) -> str:
    text = (text or "").replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()

def extract_resume_text(raw: bytes, suffix: str) -> str:
    if suffix == ".pdf":
        try:
            import PyPDF2
            reader = PyPDF2.PdfReader(BytesIO(raw))
            text = "\n\n".join((p.extract_text() or "") for p in reader.pages)
        except Exception as exc:
            raise ValueError(f"Could not parse PDF: {type(exc).__name__}.")
    elif suffix == ".docx":
        try:
            from docx import Document
            doc = Document(BytesIO(raw))
            blocks = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    blocks.append(" | ".join(c.text for c in row.cells))
            text = "\n".join(blocks)
        except Exception as exc:
            raise ValueError(f"Could not parse DOCX: {type(exc).__name__}.")
    else:
        text = raw.decode("utf-8", errors="replace")

    text = clean(text)
    if not text:
        raise ValueError("No readable text was extracted. If this is a scanned PDF, OCR is required.")
    return text
